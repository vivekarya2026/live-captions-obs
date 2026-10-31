"""
Live caption translator: English -> Hinglish/Hindi and Hindi/Hinglish -> English.

Pipeline (runs offline, on macOS and Windows):

  Qwen3-ASR (own mic capture) ---------------------+
                                                   +--> clean -> vocabulary -> translate
  LocalVocal/Whisper in OBS -> DATA/raw/speech.txt +        -> clean -> DATA/live_caption.txt
  (only until Qwen3 has been downloaded)                       (shown by the OBS caption overlay)

Rules:
  * Only English and Hindi are translated (each into the other).
  * Speech in any other language/script is dropped (almost always noise).
  * Abusive words (vocabulary/blocked_words.txt) are removed before and after
    translation.
  * Terms from vocabulary/Vocabulary_for_Live_Translations.xlsx are kept
    intact and, in English output, written with their explanation, e.g.
    "sewa (selfless service)". They are also given to Qwen3-ASR as context.

Run:  python -m livecaptions             (started by OBS; keeps running)
      python -m livecaptions --test      (translate sample lines)
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import logging
import os
import queue
import re
import shutil
import signal
import sys
import threading
import time
from pathlib import Path

from . import models, obs_bridge, paths
from .speech import NOISE_LEVELS, SPEECH_MODELS, Speech, input_devices, local_mic, save_local_mic

HERE = paths.PACKAGE
VOCAB_DIR = paths.VOCAB_DIR
VOCAB_XLSX = VOCAB_DIR / "Vocabulary_for_Live_Translations.xlsx"
HINDI_TERMS_CSV = VOCAB_DIR / "hindi_terms.csv"
BLOCKED_WORDS = VOCAB_DIR / "blocked_words.txt"

DATA = paths.DATA
RAW_DIR = paths.RAW_DIR
INPUTS = {"speech": RAW_DIR / "speech.txt"}  # written by LocalVocal (Whisper fallback)
OUTPUT = DATA / "live_caption.txt"
TRANSCRIPTS = DATA / "transcripts"
# One folder per session: Documents/Live Captions/<date>/<time> <mode>/
SESSIONS_DIR = paths.SESSIONS
HINGLISH_LEXICON = VOCAB_DIR / "hinglish_lexicon.json"
GDOCS_SCRIPT = paths.RESOURCES / "google-docs" / "LiveCaptions.gs"  # pasted into Google Docs
NO_TEXT_SOURCE = "(file output only)"
QWEN_MODEL_FILE = DATA / "speech_model.txt"   # qwen-0.6b (default) or qwen-1.7b

POLL_SECONDS = 0.05
CLEAR_AFTER_SECONDS = 5  # speech captions disappear this long after the speaker stops
FRAGMENT_JOIN_SECONDS = 2.5  # a half sentence is joined with speech that follows within this
STANDALONE_IDLE_QUIT = 180   # without OBS: stop this long after the app window was closed


def open_app_window() -> None:
    """The app window (the same page as the OBS dock), in the default browser."""
    import webbrowser
    webbrowser.open(f"http://127.0.0.1:{CONTROL_PORT}/")


def is_fragment(text: str) -> bool:
    """A short piece without an end mark ("हर पल में", "so his"): the speaker paused mid-sentence."""
    t = text.strip()
    return bool(t) and t[-1] not in ".?!।॥\"'" and len(t.split()) <= 6

# Caption direction, chosen in the OBS "Caption Language" panel.
#   auto        : English -> Hinglish and Hindi/Hinglish -> English (detected per sentence)
#   en-hinglish : speaker uses English, captions in Hinglish (Hindi in Roman letters)
#   en-hi       : speaker uses English, captions in Hindi (Devanagari)
#   hi-en       : speaker uses Hindi or Hinglish, captions in English
MODES = {
    "en-hinglish": {"label": "English → Hinglish", "hint": "speaker uses English", "whisper": "en",
                    "folder": "English to Hinglish"},
    "hi-en": {"label": "Hindi → English", "hint": "speaker uses Hindi", "whisper": "hi",
              "folder": "Hindi-Hinglish to English"},
    "auto": {"label": "Auto", "hint": "English ↔ Hinglish", "whisper": "auto",
             "folder": "Auto"},
    "en-hi": {"label": "English → हिंदी", "hint": "Hindi script", "whisper": "en",
              "folder": "English to Hindi (Devanagari)"},
}
MODE_FILE = DATA / "caption_mode.txt"
CONTROL_PORT = int(os.environ.get("LIVECAPTIONS_PORT", "8765"))  # the OBS dock opens http://127.0.0.1:8765/dock
LOCALVOCAL_INPUTS = ["Mic/Aux"]
LOCALVOCAL_FILTER = "LocalVocal Transcription"

# NLLB language codes
ENG, HIN, URD = "eng_Latn", "hin_Deva", "urd_Arab"

log = logging.getLogger("translator")

# Letters that belong to a word in Latin or Devanagari text (used as word
# boundaries, since \b breaks on Devanagari vowel signs).
WORD_CHARS = r"0-9A-Za-z\u00C0-\u024F\u0900-\u0963\u0966-\u097F"  # excludes । ॥
B_START = rf"(?<![{WORD_CHARS}])"
B_END = rf"(?![{WORD_CHARS}])"

DEVANAGARI = re.compile(r"[\u0900-\u097F]")
DEVANAGARI_LETTER = re.compile(r"[\u0900-\u0963\u0966-\u097F]")  # without । ॥
ARABIC = re.compile(r"[\u0600-\u06FF]")
LATIN = re.compile(r"[A-Za-z\u00C0-\u024F]")
NON_ASCII_LATIN = re.compile(r"[\u00C0-\u024F]")
OTHER_LETTER = re.compile(r"[^\W\d_A-Za-z\u00C0-\u024F\u0900-\u097F\u0600-\u06FF]")

# Phrases Whisper commonly invents during silence.
HALLUCINATIONS = re.compile(
    r"^(thanks? for watching|please subscribe|subscribe|subtitles? by.*|"
    r"transcribed by.*|amara\.org.*|you|bye\.?|\.+|-+)$",
    re.IGNORECASE,
)


def norm_key(s: str) -> str:
    """Lowercase, drop punctuation, collapse spaces (for matching table keys)."""
    s = re.sub(rf"[^{WORD_CHARS}\s]", " ", s.lower())
    return re.sub(r"\s+", " ", s).strip()


def phrase_pattern(phrase: str) -> str:
    """Regex for a phrase: flexible spaces/punctuation between words."""
    words = norm_key(phrase).split(" ")
    return B_START + r"[\s,.;:!?-]*".join(re.escape(w) for w in words) + B_END


def _skeleton(word: str) -> str:
    """Consonant outline of a Roman word ("sewa", "suya", "seva" -> "s"), used to
    find which vocabulary term a misheard word was meant to be."""
    w = word.lower().replace("ph", "f").replace("sh", "s").replace("x", "ks").replace("v", "w")
    return w[:1] + re.sub(r"[aeiouwyh]", "", w[1:])


def _mostly_english(text: str) -> bool:
    words = re.findall(r"[A-Za-z]+", text.lower())
    return len(words) >= 2 and sum(w in ENGLISH_WORDS for w in words) / len(words) >= 0.6


def restore_terms(out: str, terms: list[str]) -> str:
    """Put protected vocabulary terms back when the translator altered them
    ("Aatma" -> "Atma", "Avtar" -> "Avatar", "वसुधैव कुटुंबकम" -> "वसुधैव परिवारकम")."""
    for term in sorted(set(terms), key=len, reverse=True):
        if not term or term.lower() in out.lower():
            continue
        tw = term.split()
        words = list(re.finditer(r"[\w\u0900-\u097F]+", out))
        best = None
        for i in range(len(words) - len(tw) + 1):
            span = out[words[i].start():words[i + len(tw) - 1].end()]
            d = _edit_distance(span.lower(), term.lower())
            ratio = d / max(len(term), 1)
            similar = (_skeleton(span) == _skeleton(term) if term.isascii() else
                       span.split()[0] == tw[0] or ratio <= 0.34)
            if similar and ratio <= 0.5 and (best is None or d < best[0]):
                best = (d, words[i].start(), words[i + len(tw) - 1].end())
        if best:
            _, a, b = best
            out = out[:a] + term + out[b:]
    return out


def _edit_distance(a: str, b: str) -> int:
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


# --------------------------------------------------------------------------
# Vocabulary
# --------------------------------------------------------------------------
class Vocabulary:
    def __init__(self) -> None:
        self.mtimes: tuple = ()
        self.load()

    def _files(self):
        return [VOCAB_XLSX, HINDI_TERMS_CSV, BLOCKED_WORDS]

    def reload_if_changed(self) -> None:
        m = tuple(p.stat().st_mtime if p.exists() else 0 for p in self._files())
        if m != self.mtimes:
            self.load()

    def load(self) -> None:
        self.mtimes = tuple(p.stat().st_mtime if p.exists() else 0 for p in self._files())
        self.english_sub: dict[str, str] = {}    # term key -> "sewa (selfless service)"
        self.canonical: dict[str, str] = {}      # term key -> "sewa" (no explanation)
        self.hindi_to_term: dict[str, str] = {}  # Devanagari spelling -> term key
        self.term_to_hindi: dict[str, str] = {}  # term key -> preferred Devanagari
        self.hinglish_terms: dict[str, str] = {}  # Devanagari -> exact Hinglish spelling
        self.mentions: dict[str, int] = {}       # term key -> how common (orders the ASR context)
        self.last_protected: list[str] = []      # terms protected in the last translation
        self._load_hindi_terms_csv()             # hand-made spellings (full forms, variants)
        self._load_vocab()                       # the sheet: always wins
        self._finish()
        self._load_blocked()
        log.info("Loaded %d vocabulary terms, %d Hindi spellings, %d blocked words",
                 len(self.english_sub), len(self.hindi_to_term), len(self.blocked_terms))

    # ---- loading -----------------------------------------------------------
    @staticmethod
    def _display(original: str, english: str, sub: str) -> str:
        """Column C of the sheet; formula cells mean '<word> (<English>)'."""
        if not sub or sub.startswith("="):
            return f"{original} ({english})" if english and norm_key(english) != norm_key(original) \
                else (english or original)
        return sub

    def _add_roman(self, roman: str, display: str, mentions: int = 0) -> str:
        key = norm_key(roman)
        if not key:
            return key
        self.english_sub[key] = display
        self.canonical[key] = re.sub(r"\s*\(.*\)\s*$", "", display).strip() or roman
        self.mentions[key] = max(self.mentions.get(key, 0), mentions)
        return key

    def _add_hindi(self, dev: str, key: str, hinglish: str | None = None) -> None:
        dev = " ".join(dev.replace(",", " , ").split()).replace(" , ", ", ")
        self.hindi_to_term[dev] = key
        self.term_to_hindi.setdefault(key, dev)
        if hinglish:
            self.hinglish_terms[dev] = hinglish

    def _load_hindi_terms_csv(self) -> None:
        if not HINDI_TERMS_CSV.exists():
            return
        with open(HINDI_TERMS_CSV, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                term = (row.get("term") or "").strip()
                spellings = [s.strip() for s in (row.get("hindi") or "").split("|") if s.strip()]
                if term and spellings:
                    for s in spellings:
                        self._add_hindi(s, norm_key(term))

    def _load_vocab(self) -> None:
        """Both sheets of Vocabulary_for_Live_Translations.xlsx:

        vocab_live_translations   Original Word | English Translation | Substitution
        Hindi_English_Hinglish    Hindi | English | Hinglish | Substitution | Mentions
        """
        import openpyxl

        if not VOCAB_XLSX.exists():
            log.warning("Vocabulary file missing: %s", VOCAB_XLSX)
            return
        wb = openpyxl.load_workbook(VOCAB_XLSX, read_only=True)
        sheets = {ws.title.lower(): ws for ws in wb.worksheets}
        hhe = sheets.get("hindi_english_hinglish")
        dev_rows = set()
        if hhe is not None:
            for row in hhe.iter_rows(min_row=2, values_only=True):
                dev, english, roman, sub, mentions = (list(row) + [None] * 5)[:5]
                dev, english, roman = (str(x).strip() if x else "" for x in (dev, english, roman))
                if not dev or not roman:
                    continue
                sub = str(sub).strip() if sub else ""
                # English captions: Roman word + meaning (column D). Hinglish captions:
                # D's spelling when it normalises a name (सुधीक्षा -> Sudiksha), but
                # column C when D is an English sentence (तू ही निरंकार..., विश्व बंधुत्व).
                display = self._display(roman, english, sub)
                key = self._add_roman(roman, display, int(mentions or 0))
                canon = self.canonical[key]
                canon_key = norm_key(canon)
                if canon_key and canon_key != key:  # e.g. "Dhan Nirankar" -> "Dhan Nirankar Ji"
                    self._add_roman(canon, display, int(mentions or 0))
                hinglish = roman if (sub.startswith("=") or not sub or _mostly_english(canon)) else canon
                self._add_hindi(dev, key, hinglish)
                dev_rows.add(dev)
        main = sheets.get("vocab_live_translations") or wb.worksheets[0]
        for row in main.iter_rows(min_row=2, values_only=True):
            original, english, sub = (list(row) + [None, None, None])[:3]
            if not original or not str(original).strip():
                continue
            original, english = str(original).strip(), str(english or "").strip()
            sub = str(sub).strip() if sub else ""
            if DEVANAGARI_LETTER.search(original):
                if original in dev_rows:
                    continue  # already read from the Hindi_English_Hinglish sheet
                key = self._add_roman(english or original, sub if sub and not sub.startswith("=") else english)
                self._add_hindi(original, key)
                continue
            self._add_roman(original, self._display(original, english, sub))
        wb.close()

    def _finish(self) -> None:
        # Full canonical forms ("Dhan Nirankar Ji" for "dnj") are matched too, so a
        # term inside them ("Nirankar") is not explained a second time.
        for key, canon in list(self.canonical.items()):
            ck = norm_key(canon)
            if ck and ck not in self.english_sub:
                self.english_sub[ck] = self.english_sub[key]
                self.canonical[ck] = canon
        # English -> Hindi uses the Hindi form that matches the term best: as many
        # words as the term ("Dhan Nirankar Ji" -> धन निरंकार जी, not the short धनंकार;
        # "Jaikara" -> जयकारा, not the whole invocation).
        by_canon: dict[str, list[str]] = {}
        for key, canon in self.canonical.items():
            if key in self.term_to_hindi:
                by_canon.setdefault(norm_key(canon), []).append(self.term_to_hindi[key])
        for ck, devs in by_canon.items():
            n = len(ck.split())
            self.term_to_hindi[ck] = min(devs, key=lambda d: (abs(len(d.split()) - n), -len(d)))
        # English regex, longest terms first, so "satguru mata" wins over "mata".
        keys = sorted(self.english_sub, key=len, reverse=True)
        self.known_words = {w for k in keys for w in k.split()} | {"ji", "dhan"}
        self.term_patterns = [(k, phrase_pattern(k)) for k in keys]
        self.english_re = (re.compile("|".join(f"(?P<t{i}>{p})" for i, (_, p) in
                                                enumerate(self.term_patterns)), re.IGNORECASE)
                           if keys else None)
        # Devanagari regex
        hk = sorted(self.hindi_to_term, key=len, reverse=True)
        self.hindi_re = (re.compile("|".join(f"(?P<h{i}>{phrase_pattern(k)})" for i, k in enumerate(hk)))
                         if hk else None)
        self.hindi_keys = hk
        # Hinglish captions: the sheet's Hinglish spelling, else the canonical term.
        # The Hinglish column capitalises every word ("Sewa"); when the main sheet
        # has the same word, use its spelling ("sewa") so sentences read naturally.
        for dev, roman in list(self.hinglish_terms.items()):
            c = self.canonical.get(norm_key(roman))
            if c and norm_key(c) == norm_key(roman):
                self.hinglish_terms[dev] = c
        for dev, key in self.hindi_to_term.items():
            self.hinglish_terms.setdefault(dev, self.canonical.get(key, key))
        # Roman spelling -> Devanagari, for reading Hinglish input
        self.roman_to_hindi: dict[str, str] = {}
        for key, dev in self.term_to_hindi.items():
            for w in [key] + [norm_key(self.canonical.get(key, key))]:
                if " " not in w:
                    self.roman_to_hindi[w] = dev
        for dev, roman in self.hinglish_terms.items():
            k = norm_key(roman)
            if k and " " not in k:
                self.roman_to_hindi.setdefault(k, dev)
        # Near-miss correction of speech recognition ("suya" -> "sewa")
        self._roman_single = {}
        for key in self.english_sub:
            if " " not in key and len(key) >= 4:
                self._roman_single.setdefault(_skeleton(key), []).append(key)
        self._dev_single = [d for d in self.hindi_to_term if " " not in d and len(d) >= 4]
        self.asr_context = self.asr_context_for("auto")

    def asr_context_for(self, language: str | None, limit: int = 700) -> str:
        """Vocabulary for Qwen3-ASR's context prompt, most common terms first:
        Roman spellings for English speech, Devanagari for Hindi, both for auto.
        ~700 characters: the full list made each sentence ~0.3 s slower for no
        accuracy gain; a much shorter one made the model sometimes repeat it."""
        order = sorted(self.canonical, key=lambda k: -self.mentions.get(k, 0))
        roman = list(dict.fromkeys(self.canonical[k] for k in order if len(self.canonical[k]) <= 60))
        dev = sorted(self.hindi_to_term, key=lambda d: -self.mentions.get(self.hindi_to_term[d], 0))
        dev = list(dict.fromkeys(d for d in dev if len(d) <= 60))
        if language == "en":
            text = "Nirankari Mission satsang. Spell these words exactly: " + ", ".join(roman) + "."
        elif language == "hi":
            text = "निरंकारी मिशन सत्संग। इन शब्दों को ठीक ऐसे ही लिखें: " + ", ".join(dev) + "।"
        else:
            text = ("Nirankari Mission satsang vocabulary: " + ", ".join(roman[:60]) + ". "
                    + ", ".join(dev[:80]) + "।")
        return text[:limit]

    def fix_recognition(self, text: str, lang: str) -> str:
        """Snap near-miss spellings from speech recognition to vocabulary terms
        (only words that aren't real English/Hindi words)."""
        if lang == "en" and self._roman_single:
            def fix(m: re.Match) -> str:
                w = m.group(0)
                lw = w.lower()
                if len(lw) < 4 or lw in ENGLISH_WORDS or lw in self.english_sub:
                    return w
                best = None
                for key in self._roman_single.get(_skeleton(lw), []):
                    d = _edit_distance(lw, key)
                    if d <= max(2, len(key) // 3) and (best is None or d < best[0]):
                        best = (d, key)
                if not best:
                    return w
                out = self.canonical.get(best[1], best[1]).split(" (")[0]
                return out[:1].upper() + out[1:] if w[:1].isupper() else out
            return re.sub(r"[A-Za-z]+", fix, text)
        if lang == "hi" and self._dev_single:
            def fix_dev(m: re.Match) -> str:
                w = m.group(0)
                if len(w) < 4 or w in self.hindi_to_term or w in self._known_hindi:
                    return w
                best = None
                for d in self._dev_single:
                    if abs(len(d) - len(w)) > 1:
                        continue
                    dist = _edit_distance(w, d)
                    if dist <= (1 if len(d) < 7 else 2) and (best is None or dist < best[0]):
                        best = (dist, d)
                return best[1] if best else w
            return re.sub(r"[\u0900-\u0963\u0966-\u097F]+", fix_dev, text)
        return text

    def is_context_echo(self, text: str) -> bool:
        """True if the transcript is just a list of vocabulary terms, which
        Qwen3-ASR sometimes outputs (from its context prompt) on noise/silence."""
        parts = [p.strip(" .।") for p in re.split(r"[,،]", text) if p.strip(" .।")]
        if len(parts) < 4:
            return False
        known = sum(1 for p in parts if p in self.hindi_to_term or norm_key(p) in self.english_sub)
        return known / len(parts) >= 0.6

    _known_hindi: set = set()  # set by CaptionPipeline (lexicon words that must not be "fixed")

    def _load_blocked(self) -> None:
        self.blocked_terms: list[str] = []
        parts = []
        if BLOCKED_WORDS.exists():
            for line in BLOCKED_WORDS.read_text(encoding="utf-8").splitlines():
                w = line.strip()
                if not w or w.startswith("#"):
                    continue
                self.blocked_terms.append(w)
                stem = w.rstrip("*")
                body = r"\s+".join(re.escape(x) for x in stem.split())
                tail = rf"[{WORD_CHARS}]*" if w.endswith("*") else ""
                parts.append(B_START + body + tail + B_END)
        self.blocked_re = re.compile("|".join(parts), re.IGNORECASE) if parts else None

    # ---- operations -------------------------------------------------------
    def remove_abusive(self, text: str) -> str:
        if not self.blocked_re:
            return text
        text = self.blocked_re.sub("", text)
        text = re.sub(r"\s+([,.;:!?।])", r"\1", text)
        text = re.sub(r"([,;:])\1+", r"\1", text)
        return re.sub(r"\s{2,}", " ", text).strip(" ,;:-")

    def _term_key_for_match(self, m: re.Match) -> str:
        return self.term_patterns[int(m.lastgroup[1:])][0]

    def apply_english(self, text: str) -> str:
        """English output: 'sewa' -> 'sewa (selfless service)' (once, no repeats)."""
        if not self.english_re:
            return text

        def repl(m: re.Match) -> str:
            sub = self.english_sub[self._term_key_for_match(m)]
            before = text[:m.start()].rstrip()
            sentence_start = not before or before[-1] in ".!?:\"'“"
            # Capital only at the start of a sentence for words the sheet writes in
            # lowercase ("do sewa", not "do Sewa"); names keep the sheet's capitals.
            if sentence_start and sub[:1].islower():
                sub = sub[:1].upper() + sub[1:]
            return sub

        # Skip terms that already carry an explanation "(...)" after them.
        out, pos = [], 0
        for m in self.english_re.finditer(text):
            if re.match(r"\s*\(", text[m.end():]):
                continue
            out.append(text[pos:m.start()])
            out.append(repl(m))
            pos = m.end()
        out.append(text[pos:])
        return "".join(out)

    def protect_in_english_source(self, text: str) -> str:
        """English -> Hindi: swap vocabulary terms for their Hindi spelling first,
        so the translator keeps them instead of translating them literally."""
        if not self.english_re:
            return text

        def repl(m: re.Match) -> str:
            key = self._term_key_for_match(m)
            canon = norm_key(self.canonical.get(key, key))
            out = self.term_to_hindi.get(key) or self.term_to_hindi.get(canon) or self.canonical.get(key, m.group(0))
            self.last_protected.append(out)
            return out

        self.last_protected = []
        return self.english_re.sub(repl, text)

    def protect_in_hindi_source(self, text: str) -> str:
        """Hindi -> English: swap Devanagari vocabulary terms for the romanized
        term in quotes (e.g. सेवा -> "Sewa"). The translator copies quoted,
        capitalized terms through unchanged; quotes are removed afterwards."""
        if not self.hindi_re:
            return text

        def repl(m: re.Match) -> str:
            term = self.hindi_to_term[self.hindi_keys[int(m.lastgroup[1:])]]
            # Everyday phrases, not the title: "माता पिता" / "माता-पिता" = parents
            if m.group(0) == "माता" and re.match(r"[\s-]*पिता", text[m.end():]):
                return m.group(0)
            term = self.canonical.get(norm_key(term), term)  # e.g. dnj -> Dhan Nirankar Ji
            term = term[:1].upper() + term[1:]
            if not _mostly_english(term):
                self.last_protected.append(term)
            return f'"{term}"'

        self.last_protected = []
        text = self.hindi_re.sub(repl, text)
        # Keep runs of names/titles together, e.g.
        # "Satguru Mata" "Sudiksha" जी "Maharaj" -> "Satguru Mata Sudiksha Ji Maharaj"
        text = re.sub(r'"\s+जी' + B_END, ' Ji"', text)
        prev = None
        while prev != text:
            prev = text
            text = re.sub(r'"([^"]+)"\s+"([^"]+)"', r'"\1 \2"', text)
        return text

    @staticmethod
    def unquote_terms(text: str) -> str:
        return re.sub(r'["“”]([^"“”]{1,80})["“”]', r"\1", text)

    def to_hinglish(self, hindi: str) -> str:
        """Devanagari -> Hinglish, with vocabulary terms in the sheet's spelling."""
        from .hinglish import to_hinglish
        out = to_hinglish(hindi, self.hinglish_terms)
        # Keep the sheet's spelling/capitalisation for names and titles (Satguru,
        # Nirankar, Sudhiksha -> Sudiksha), but never swap in an English meaning
        # ("Tu hi Nirankar..." stays Hinglish, not "Only YOU are...").
        if self.english_re:
            def keep(m: re.Match) -> str:
                canon = self.canonical.get(self._term_key_for_match(m), "")
                if canon[:1].isupper() and not _mostly_english(canon) and (
                        norm_key(canon) == norm_key(m.group(0))
                        or _edit_distance(norm_key(canon), norm_key(m.group(0))) <= max(2, len(canon) // 4)):
                    return canon
                return m.group(0)
            out = self.english_re.sub(keep, out)
        return out


# --------------------------------------------------------------------------
# Translation model (Meta NLLB-200 1.3B, int8, via CTranslate2)
# --------------------------------------------------------------------------
class Translator:
    def __init__(self) -> None:
        import ctranslate2
        import sentencepiece as spm

        t0 = time.time()
        model_dir = models.model_dir("nllb")
        threads = max(2, min(6, (os.cpu_count() or 4) // 2))  # 6 was fastest next to OBS on an M4 Pro
        self.model = ctranslate2.Translator(str(model_dir), device="cpu", compute_type="int8",
                                            intra_threads=threads)
        self.sp = spm.SentencePieceProcessor(model_file=str(model_dir / "sentencepiece.bpe.model"))
        self.cache: dict[tuple, str] = {}
        log.info("Translation model loaded in %.1fs", time.time() - t0)

    def translate(self, text: str, src: str, tgt: str) -> str:
        # Translate sentence by sentence: NLLB tends to drop later sentences otherwise.
        sentences = [s.strip() for s in re.split(r"(?<=[.!?।|])\s+", text) if s.strip()]
        todo = [s for s in sentences if (s, src, tgt) not in self.cache]
        if todo:
            batch = [[src] + self.sp.encode(s, out_type=str) + ["</s>"] for s in todo]
            # beam_size=1 (greedy): ~40% faster than beam 4 with near-identical output
            res = self.model.translate_batch(batch, target_prefix=[[tgt]] * len(batch), beam_size=1,
                                             max_decoding_length=256, repetition_penalty=1.1)
            for s, r in zip(todo, res):
                self.cache[(s, src, tgt)] = self.sp.decode(r.hypotheses[0][1:]).strip()
            if len(self.cache) > 2000:
                self.cache.clear()
        return " ".join(self.cache[(s, src, tgt)] for s in sentences)


# --------------------------------------------------------------------------
# Caption processing
# --------------------------------------------------------------------------
def _load_english_words() -> set[str]:
    words: set[str] = set()
    # Word list shipped with the app (from the macOS dictionary), so Windows works the same.
    for p in (paths.RESOURCES / "english_words.txt", Path("/usr/share/dict/words")):
        if p.exists():
            for w in p.read_text(encoding="utf-8", errors="ignore").split():
                words.add(w.lower())
            break
    # Common words/contractions the system dictionary lacks.
    words.update("i'm you're we're they're it's don't can't won't let's i've we'll "
                 "isn't aren't didn't doesn't that's there's okay ok hello hi".split())
    return words


ENGLISH_WORDS = _load_english_words()


def looks_like_english(text: str, vocab: "Vocabulary", lenient: bool = False) -> bool:
    """True if most words are real English words or vocabulary terms.
    Filters out Whisper noise such as "Erne Dopenbahn" during silence."""
    if not ENGLISH_WORDS:
        return True
    words = re.findall(r"[A-Za-z']+", text.lower())
    if not words:
        return False
    known = sum(1 for w in words
                if w in ENGLISH_WORDS or w.strip("'") in ENGLISH_WORDS
                or re.sub(r"'s$", "", w) in ENGLISH_WORDS  # "mom's"
                or w.rstrip("s") in ENGLISH_WORDS or w in vocab.known_words)
    return known / len(words) >= (0.6 if lenient else 0.75)


def _load_hindi_frequencies() -> dict[str, int]:
    """Common spoken Hindi words with how often they occur (from subtitles)."""
    freq: dict[str, int] = {}
    p = paths.RESOURCES / "hindi_words.txt"
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            parts = line.split()
            if len(parts) == 2 and not line.startswith("#"):
                freq[parts[0]] = int(parts[1])
    return freq


HINDI_FREQ = _load_hindi_frequencies()


def hindi_skeleton(word: str) -> str:
    """Spelling-insensitive form of a Hindi word: the variations a speech model
    gets wrong (ी/ि, ू/ु, ँ/ं, ृ/्र, nukta, half letters) all map to the same
    skeleton, e.g. जाग्रुक and जागरूक -> जागरुक, जाग्रति and जागृति -> जागरति."""
    w = word.replace("\u0943", "\u0930")          # ृ -> र
    w = w.replace("\u094D", "").replace("\u093C", "")  # half letters, nukta
    w = w.replace("\u0901", "\u0902")              # ँ -> ं
    return w.translate({0x0940: 0x093F, 0x0942: 0x0941, 0x0908: 0x0907, 0x090A: 0x0909})


SHORT_HINDI_OK = {"नमस्ते", "धन्यवाद", "शुक्रिया", "जयकारा", "जैकारा"}

# English words written in Hindi letters ("ऑपरेशन", "इंसिडेंट", "एक्सेप्ट"):
# sounds Hindi words almost never have.
ENGLISH_IN_DEVANAGARI = re.compile(
    "[\u0911\u0949]|(िंग|ेंट|ेंस|मेंट|टर|्ड)$|^(एक्स|इं|फ्)")

# Words of prayer the speech model mishears in satsang (seen in live tests).
SPOKEN_FIXES = {
    **{w: "बख्शना" for w in ("बक्षना", "भक्षना", "वक्षना", "बक्शना", "बख्षना", "बखशना")},
    **{w: "बख्शो" for w in ("बक्षो", "भक्षो", "बक्शो")},
    **{w: "बख्श" for w in ("बक्ष", "बक्श")},
    **{w: "हौसला" for w in ("हसला", "हासला", "होसला", "हौंसला")},
    "जाग्रुप": "जागरूक", "जाग्रुक": "जागरूक", "खुशामत": "खुशामद", "सद्भुद्धि": "सद्बुद्धि",
    "जाग्रुती": "जागृति", "जाग्रूती": "जागृति", "संगमगम": "समागम", "समगम": "समागम",
    "स्तिर": "स्थिर", "स्तीर": "स्थिर", "संदिश": "संदेश", "सहूलत": "सहूलियत",
}

# Greetings / phrases of the satsang the speech model mishears as a whole.
_W = r"(?<![\u0900-\u097F])"  # start of a Hindi word
SPOKEN_PHRASE_FIXES = [
    (_W + r"धन\s+नि(रा|र)कार(\s+जी)?", r"धन निरंकार\2"),           # "धन निराकार जी"
    (_W + r"(सत्संगे|साद संगत|सद संगत|साध संगते)\s+जी", "साध संगत जी"),  # greeting to the sangat
]

# Hindi words the translator misreads; replaced only in what it is given.
TRANSLATION_HINTS = [(_W + r"दात(?![\u0900-\u097F])", "उपहार")]  # "ज्ञान की दात" = gift, not teeth


def looks_like_hindi(text: str, pipeline: "CaptionPipeline") -> bool:
    """True if most Devanagari words are real Hindi words or vocabulary terms.
    Filters out what IndicWhisper makes of background noise, e.g.
    "िवार और भाविविविविवध बनाएक भी"."""
    words = re.findall(r"[\u0900-\u0963\u0966-\u097F]+", text)
    if not words:
        return False
    if re.match(r"[\u093E-\u094D\u0900-\u0903]", words[0]):  # starts with a stray vowel sign
        return False
    known = pipeline.hindi_words

    def is_known(w: str) -> bool:
        return (w in known or w in pipeline.vocab.hindi_to_term
                or hindi_skeleton(w) in pipeline.hindi_skeletons
                or (pipeline.lenient and ENGLISH_IN_DEVANAGARI.search(w) is not None))

    if pipeline.lenient:
        # Qwen3 only hears audio the voice detector passed and seldom invents
        # words, so keep anything that is mostly real Hindi (also short replies).
        ok = sum(1 for w in words if is_known(w))
        return ok / len(words) >= (0.5 if len(words) <= 3 else 0.6)
    if len(words) == 1 and words[0] not in pipeline.vocab.hindi_to_term and words[0] not in SHORT_HINDI_OK:
        return False  # a lone word is almost always noise ("एल", "हाँ" from the room)
    ok = sum(1 for w in words if is_known(w))
    return ok / len(words) >= (0.75 if len(words) <= 3 else 0.6)


def detect_language(text: str) -> str | None:
    """'hi', 'ur' (Hindi speech written in Urdu script), 'en', 'other' or None."""
    dev, ara, lat = len(DEVANAGARI.findall(text)), len(ARABIC.findall(text)), len(LATIN.findall(text))
    other = len(OTHER_LETTER.findall(text))
    total = dev + ara + lat + other
    if total < 2:
        return None
    if dev / total >= 0.5:
        return "hi"
    if ara / total >= 0.5:
        return "ur"
    if lat / total >= 0.6:
        # Lots of accented letters (e.g. Turkish/French) -> not English
        if len(NON_ASCII_LATIN.findall(text)) / max(lat, 1) > 0.05:
            return "other"
        return "en"
    return "other"


class CaptionPipeline:
    def __init__(self, translator: Translator | None) -> None:
        self.vocab = Vocabulary()
        self.translator = translator
        self.mode = read_mode()
        self.obs_status = ""
        from .hinglish import HinglishLexicon
        self.lexicon = HinglishLexicon(HINGLISH_LEXICON, self.vocab.roman_to_hindi)
        # Known Hindi words (for dropping noise transcriptions)
        from .hinglish import COMMON
        self.hindi_words = set(self.lexicon.exact.values()) | set(self.lexicon.loose.values()) | set(COMMON)
        for spelling in self.vocab.hindi_to_term:
            self.hindi_words.update(spelling.split())
        self.hindi_words.update(HINDI_FREQ)
        for fixed in SPOKEN_FIXES.values():
            self.hindi_words.update(fixed.split())
        self.vocab._known_hindi = self.hindi_words
        # skeleton -> correct spelling (vocabulary first, then the most common word)
        self.hindi_skeletons: dict[str, str] = {}
        ranked = sorted(self.hindi_words, key=lambda w: -HINDI_FREQ.get(w, 0))
        vocab_words = [w for s in self.vocab.hindi_to_term for w in s.split()]
        for w in vocab_words + ranked:
            self.hindi_skeletons.setdefault(hindi_skeleton(w), w)
        self.lenient = False  # set while processing Qwen3 speech
        self.last: dict = {}  # details of the last processed line (for session logs)
        self.recorder = SessionRecorder()

    def fix_hindi_spelling(self, text: str) -> str:
        """Unknown words the speech model misspelt -> the common spelling with
        the same sound (जाग्रुक -> जागरूक, जाग्रति -> जागृति)."""
        for wrong, right in SPOKEN_PHRASE_FIXES:
            text = re.sub(wrong, right, text)

        def fix(m: re.Match) -> str:
            w = m.group(0)
            if w in SPOKEN_FIXES:
                return SPOKEN_FIXES[w]
            if len(w) < 3 or w in self.hindi_words or w in self.vocab.hindi_to_term:
                return w
            return self.hindi_skeletons.get(hindi_skeleton(w), w)
        return re.sub(r"[\u0900-\u0963\u0970-\u097F]+", fix, text)

    def _hindi_out(self, hindi: str) -> str:
        """Hindi target: Hinglish (Roman) unless the mode asks for Devanagari."""
        return hindi if self.mode == "en-hi" else self.vocab.to_hinglish(hindi)

    def _hindi_to_english(self, text: str, src_lang: str) -> str:
        src_text = self.vocab.protect_in_hindi_source(text) if src_lang == "hi" else text
        for wrong, right in TRANSLATION_HINTS:
            src_text = re.sub(wrong, right, src_text)
        if src_lang == "hi" and not DEVANAGARI_LETTER.search(src_text):
            out = src_text  # only vocabulary terms: nothing left to translate
        else:
            out = self.translator.translate(src_text, HIN if src_lang == "hi" else URD, ENG) \
                if self.translator else src_text
        if src_text != text:
            out = self.vocab.unquote_terms(out)
            if out.count('"') % 2:  # a protected term's quote the translator left unpaired
                out = out.replace('"', "")
            out = restore_terms(out, self.vocab.last_protected)  # "Avatar" -> "Avtar"
        out = re.sub(r"(\w) ['’](s|t|m|re|ve|ll|d)\b", r"\1'\2", out)  # "Let 's" -> "Let's"
        return self.vocab.apply_english(out.replace("।", ".").replace("॥", "."))

    def process(self, raw: str, source: str) -> str | None:
        """Returns the caption to show, or None to keep the previous caption.
        Details (language, clean transcript) are left in self.last."""
        self.last = {}
        text = " ".join(raw.split()).strip(" -–—")
        # Sound tags Whisper adds, e.g. *music*, [Music], (applause), ♪
        text = re.sub(r"\*[^*]{1,30}\*|\[[^\]]{1,30}\]|\([^)]{0,30}(music|applause|laugh|noise|silence|"
                      r"inaudible|संगीत)[^)]{0,30}\)|[♪♫]+", "", text, flags=re.IGNORECASE).strip(" -–—")
        if not text or HALLUCINATIONS.match(text):
            return None
        if self.vocab.is_context_echo(text):
            log.info("Dropped (speech model repeated its vocabulary list): %s", text[:80])
            return None
        self.vocab.reload_if_changed()
        text = self.vocab.remove_abusive(text)
        lang = detect_language(text)
        if lang is None:
            return None
        if source == "speech" and lang in ("en", "hi"):
            fixed = text
            if lang == "hi":
                for wrong, right in SPOKEN_PHRASE_FIXES:
                    fixed = re.sub(wrong, right, fixed)
            fixed = self.vocab.fix_recognition(fixed, lang)  # "do suya" -> "do sewa"
            if lang == "hi":
                fixed = self.fix_hindi_spelling(fixed)
            if fixed != text:
                log.info("Vocabulary fix: %s  ->  %s", text, fixed)
                text = fixed

        from .hinglish import is_hinglish, to_devanagari
        if lang == "en" and is_hinglish(text, ENGLISH_WORDS):
            lang = "hinglish"

        if lang == "en" and source != "ocr" and not looks_like_english(text, self.vocab, self.lenient):
            log.info("Dropped (not real English): %s", text)
            return None
        if lang == "hi" and source != "ocr" and not looks_like_hindi(text, self):
            log.info("Dropped (not real Hindi): %s", text)
            return None

        if lang == "hinglish":
            # Roman Hindi -> Devanagari so the translator understands it
            dev = to_devanagari(text, self.lexicon, ENGLISH_WORDS, terms=self.vocab.roman_to_hindi)
            if self.mode == "en-hi":
                out = dev
            elif self.mode == "en-hinglish":
                out = self.vocab.to_hinglish(dev)  # normalise spelling and vocabulary terms
            else:
                out = self._hindi_to_english(dev, "hi")
        elif lang == "en" and self.mode == "hi-en":
            # Captions should be English and this already is: show it (with vocabulary).
            out = self.vocab.apply_english(text)
        elif lang == "en":
            src_text = self.vocab.protect_in_english_source(text)
            hindi = self.translator.translate(src_text, ENG, HIN) if self.translator else src_text
            hindi = restore_terms(hindi, self.vocab.last_protected)  # keep the sheet's Hindi spelling
            out = self._hindi_out(hindi)
        elif lang == "hi" and self.mode in ("en-hi", "en-hinglish"):
            out = self._hindi_out(text)  # captions should be Hindi and this already is
        elif lang in ("hi", "ur"):
            out = self._hindi_to_english(text, lang)
        else:
            if source != "ocr":
                log.info("Dropped (not English/Hindi): %s", text)
                return None
            out = text  # e.g. Punjabi OCR: show as-is

        out = self.vocab.remove_abusive(out)
        self.last = {"lang": lang, "transcript": text, "caption": out}
        return out or None


def write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


def log_transcript(source: str, original: str, caption: str) -> None:
    TRANSCRIPTS.mkdir(parents=True, exist_ok=True)
    now = dt.datetime.now()
    with open(TRANSCRIPTS / f"{now:%Y-%m-%d}.txt", "a", encoding="utf-8") as f:
        f.write(f"[{now:%H:%M:%S}] ({source}) {original}\n           -> {caption}\n")


LANG_NAMES = {"en": "English", "hi": "Hindi", "hinglish": "Hinglish", "ur": "Hindi (Urdu script)",
              "other": "Other"}


class SessionRecorder:
    """Saves every caption of a session into its own folder:

    Sessions/2026-10-29/17-45 English to Hinglish/
        transcript.txt    what was said (as transcribed)
        translation.txt   the captions that were shown
        bilingual.txt     both, line by line, with time and language
        captions.srt      subtitle file (open with the recording in a video player)
        session.csv       everything, for Excel

    A new folder starts when the translator starts, when the caption language
    is switched, or after 30 minutes of silence.
    """
    IDLE_NEW_SESSION = 30 * 60

    def __init__(self, root: Path = SESSIONS_DIR) -> None:
        self.root = root
        self.dir: Path | None = None
        self.mode = None
        self.start = None
        self.last_time = 0.0
        self.count = 0

    def _open(self, mode: str) -> None:
        now = dt.datetime.now()
        label = MODES[mode].get("folder", mode)
        self.dir = self.root / f"{now:%Y-%m-%d}" / f"{now:%H-%M-%S} {label}"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.mode, self.start, self.count = mode, now, 0
        header = f"Session {now:%Y-%m-%d %H:%M:%S} | {MODES[mode]['label']}\n\n"
        for name in ("transcript.txt", "translation.txt", "bilingual.txt"):
            (self.dir / name).write_text(header, encoding="utf-8")
        with open(self.dir / "session.csv", "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerow(["time", "source", "language", "transcript", "translation"])
        (self.dir / "captions.srt").write_text("", encoding="utf-8")
        log.info("Saving session to %s", self.dir)

    @staticmethod
    def _srt_time(seconds: float) -> str:
        ms = int(round(seconds * 1000))
        h, ms = divmod(ms, 3600_000)
        m, ms = divmod(ms, 60_000)
        s, ms = divmod(ms, 1000)
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"

    def add(self, mode: str, source: str, lang: str, transcript: str, caption: str) -> None:
        t = time.time()
        if self.dir is None or mode != self.mode or t - self.last_time > self.IDLE_NEW_SESSION:
            self._open(mode)
        self.last_time = t
        self.count += 1
        now = dt.datetime.now()
        stamp = f"{now:%H:%M:%S}"
        lang_name = LANG_NAMES.get(lang, lang)
        d = self.dir
        with open(d / "transcript.txt", "a", encoding="utf-8") as f:
            f.write(f"[{stamp}] {transcript}\n")
        with open(d / "translation.txt", "a", encoding="utf-8") as f:
            f.write(f"[{stamp}] {caption}\n")
        with open(d / "bilingual.txt", "a", encoding="utf-8") as f:
            f.write(f"[{stamp}] ({lang_name}, {source}) {transcript}\n           → {caption}\n\n")
        with open(d / "session.csv", "a", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerow([stamp, source, lang_name, transcript, caption])
        start = (now - self.start).total_seconds()
        with open(d / "captions.srt", "a", encoding="utf-8") as f:
            f.write(f"{self.count}\n{self._srt_time(max(0, start - 0.5))} --> "
                    f"{self._srt_time(start + 4)}\n{caption}\n\n")


def read_mode() -> str:
    try:
        m = MODE_FILE.read_text().strip()
        return m if m in MODES else "auto"
    except FileNotFoundError:
        return "auto"


NOISE_FILE = DATA / "noise_level.txt"
CAPTIONS_ON_FILE = DATA / "captions_on.txt"


def read_noise_level() -> str:
    try:
        n = NOISE_FILE.read_text().strip()
        return n if n in NOISE_LEVELS else "medium"
    except FileNotFoundError:
        return "medium"


def captions_on() -> bool:
    try:
        return CAPTIONS_ON_FILE.read_text().strip() != "off"
    except FileNotFoundError:
        return True


class App:
    """Everything the dock can change, and the background syncing with OBS."""

    def __init__(self, pipeline: CaptionPipeline, setup: models.Setup, with_obs: bool = False) -> None:
        self.pipeline = pipeline
        self.setup = setup
        self.speech = Speech(setup)
        self.with_obs = with_obs
        self.gdocs = None          # GoogleDocs link (set in run())
        self.quit = False          # "Quit" in the app window
        self.last_client = time.time()  # last time the app window / OBS dock asked for state
        self.status = ""
        self._sync_lock = threading.Lock()
        self._sync_again = False

    def configure_speech(self) -> None:
        lang_code = MODES[self.pipeline.mode]["whisper"]
        self.speech.configure(lang_code, self.pipeline.vocab.asr_context_for(lang_code),
                              read_noise_level(), captions_on())

    def sync(self, wait: bool = False) -> None:
        """Push current settings to the speech engine and OBS (in the background).
        Repeated clicks collapse into one sync with the latest settings."""
        self.configure_speech()
        self.status = "switching"

        def work():
            if not self._sync_lock.acquire(blocking=False):
                self._sync_again = True
                return
            try:
                while True:
                    self._sync_again = False
                    self.configure_speech()
                    try:
                        msg = self.speech.sync_obs()
                    except Exception as e:  # never leave the dock stuck on "switching"
                        log.exception("Could not update OBS")
                        msg = f"could not update OBS ({e.__class__.__name__})"
                    if not self._sync_again:
                        break
                self.status = msg
                log.info("Caption mode: %s (%s)", MODES[self.pipeline.mode]["label"], msg)
            finally:
                self._sync_lock.release()

        if wait:
            work()
        else:
            threading.Thread(target=work, daemon=True).start()

    def set_mode(self, mode: str) -> None:
        obs_bridge.write_atomic(MODE_FILE, mode)
        self.pipeline.mode = mode
        self.sync()

    def mic(self, payload: dict, reloads: bool) -> None:
        if obs_bridge.state() is None:  # no OBS: Live Captions listens to the mic itself
            changes = {}
            if "input_settings" in payload:
                changes["name"] = payload["input_settings"].get("device_id") or None
            for k in ("volume", "muted"):
                if k in payload:
                    changes[k] = payload[k]
            save_local_mic(**changes)
            self.speech.follow_obs_microphone()
            self.status = "saved"
            return

        def work():
            with obs_bridge.lock:
                ok = obs_bridge.state() is not None and obs_bridge.send(
                    payload, reloads=reloads and self.speech.engine == "whisper")
            self.speech.follow_obs_microphone()
            self.status = "saved" if ok else "OBS is not running"
        threading.Thread(target=work, daemon=True).start()

    def state(self) -> dict:
        obs = obs_bridge.state() is not None
        if obs:
            mic = obs_bridge.mic_state()
        else:
            lm = local_mic()
            mic = {"device_id": lm["name"] or "", "devices": input_devices(),
                   "volume": lm["volume"], "muted": lm["muted"]}
        sp = self.speech
        p = self.pipeline
        return {
            "running": True,
            "obs": obs,
            "standalone": not self.with_obs and not obs,
            "gdocs": self.gdocs.status() if self.gdocs else None,
            "engine": sp.engine,
            "model": (SPEECH_MODELS.get(sp.qwen.model_key, "Qwen3") if sp.engine == "qwen"
                      else "Whisper (until Qwen3 is ready)"),
            "speech_model": sp.model_key, "speech_status": sp.status, "speech_loading": sp.loading,
            "speech_models": {k: {"label": v, "ready": models.model_ready(k)} for k, v in SPEECH_MODELS.items()},
            "setup": self.setup.status(),
            "translator_ready": p.translator is not None,
            "captions_on": captions_on(),
            "clear_after": CLEAR_AFTER_SECONDS,
            "mode": p.mode, "status": self.status,
            "modes": {k: v["label"] for k, v in MODES.items()},
            "hints": {k: v["hint"] for k, v in MODES.items()},
            "mic": {"device_id": mic.get("device_id", ""), "devices": mic.get("devices", []),
                    "volume": mic.get("volume", 1.0), "muted": mic.get("muted", False),
                    "noise": read_noise_level()},
            "session": str(p.recorder.dir or ""), "sessions_root": str(SESSIONS_DIR),
            "lines": p.recorder.count,
            "caption": OUTPUT.read_text(encoding="utf-8") if OUTPUT.exists() else "",
        }


def start_control_server(app: App) -> bool:
    """Small local web server for the OBS 'Live Captions' dock. False if the
    port is taken (another translator is already running)."""
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from urllib.parse import parse_qs, urlparse

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def _send(self, code: int, body: bytes, ctype: str) -> None:
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)

        def _json(self, obj) -> None:
            self._send(200, json.dumps(obj).encode(), "application/json")

        def do_GET(self):
            # Only our own window / the OBS dock may use this server: refuse other
            # websites (they could e.g. send the captions to their own "Google Doc").
            host = (self.headers.get("Host") or "").split(":")[0]
            if host not in ("127.0.0.1", "localhost") or \
                    self.headers.get("Sec-Fetch-Site") in ("cross-site", "same-site"):
                return self._send(403, b"forbidden", "text/plain")
            url = urlparse(self.path)
            q = {k: v[0] for k, v in parse_qs(url.query).items()}
            busy = app.status == "switching"
            p = app.pipeline
            if url.path in ("/", "/panel", "/dock"):
                return self._send(200, (HERE / "dock.html").read_bytes(), "text/html; charset=utf-8")
            if url.path == "/session/new":
                p.recorder.dir = None  # next caption opens a new folder
                return self._json({"ok": True})
            if url.path == "/session/open":
                paths.open_folder(p.recorder.dir or SESSIONS_DIR)
                return self._json({"ok": True})
            if url.path == "/captions":
                on = q.get("on") == "1"
                write_atomic(CAPTIONS_ON_FILE, "on" if on else "off")
                if not on:
                    write_atomic(OUTPUT, "")  # clear the screen now
                app.sync()
                return self._json(app.state())
            if url.path == "/mic":
                if "device" in q and not busy:
                    app.mic({"input_settings": {"device_id": q["device"]}}, reloads=True)
                elif "volume" in q:
                    app.mic({"volume": max(0.0, min(1.0, float(q["volume"])))}, reloads=False)
                elif "muted" in q:
                    app.mic({"muted": q["muted"] == "1"}, reloads=False)
                elif q.get("noise") in NOISE_LEVELS:
                    write_atomic(NOISE_FILE, q["noise"])
                    app.sync()
                return self._json(app.state())
            if url.path == "/speech":
                if q.get("model") in SPEECH_MODELS:
                    app.speech.choose_model(q["model"], on_started=app.sync)
                return self._json(app.state())
            if url.path == "/setup/retry":
                app.setup.start()
                return self._json(app.state())
            if url.path in ("/mode", "/state"):
                app.last_client = time.time()
                mode = q.get("set")
                if mode in MODES and mode != p.mode:
                    app.set_mode(mode)
                return self._json(app.state())
            if url.path == "/gdocs":
                g = app.gdocs
                if "link" in q:
                    err = g.connect(q["link"])
                    return self._json({**app.state(), "gdocs_error": err})
                if q.get("disconnect") == "1":
                    g.disconnect()
                if q.get("write") in ("0", "1"):
                    g.command(writing=q["write"] == "1")
                if q.get("time") in ("0", "1"):
                    g.command(show_time=q["time"] == "1")
                return self._json(app.state())
            if url.path == "/gdocs/script":
                return self._send(200, GDOCS_SCRIPT.read_bytes(), "text/plain; charset=utf-8")
            if url.path == "/quit":
                app.quit = True
                return self._json({"ok": True})
            self._send(404, b"not found", "text/plain")

    try:
        server = ThreadingHTTPServer(("127.0.0.1", CONTROL_PORT), Handler)
    except OSError as e:
        log.warning("Port %d busy (%s): another Live Captions translator is running", CONTROL_PORT, e)
        return False
    threading.Thread(target=server.serve_forever, daemon=True).start()
    log.info("Live Captions dock: http://127.0.0.1:%d/dock", CONTROL_PORT)
    return True


def install_overlay() -> None:
    """The OBS browser source shows DATA/overlay.html, which polls live_caption.txt next to it."""
    src = paths.RESOURCES / "overlay.html"
    dst = DATA / "overlay.html"
    if src.exists() and (not dst.exists() or dst.read_bytes() != src.read_bytes()):
        shutil.copyfile(src, dst)


def run(with_obs: bool = False, open_window: bool = False) -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    install_overlay()
    setup = models.Setup()
    pipeline = CaptionPipeline(Translator() if models.model_ready("nllb") else None)
    app = App(pipeline, setup, with_obs=with_obs)
    if not start_control_server(app):
        if open_window:  # already running (e.g. started by OBS): just show its window
            open_app_window()
        return
    paths.PID_FILE.write_text(str(os.getpid()))
    from .gdocs import GoogleDocs
    app.gdocs = GoogleDocs()
    if not with_obs:
        from . import obs_setup
        obs_setup.finish_pending()  # the installer ran while OBS was open
    if open_window:
        open_app_window()

    def nllb_ready():
        pipeline.translator = Translator()

    setup.when_ready("nllb", nllb_ready)
    setup.when_ready("runtime", lambda: app.speech.start_qwen_if_ready(app.sync))
    for key in SPEECH_MODELS:
        setup.when_ready(key, lambda: app.speech.start_qwen_if_ready(app.sync))
    setup.start()                                    # downloads whatever is missing
    app.speech.start_qwen_if_ready(app.sync)         # already downloaded: start now
    app.sync()                                       # LocalVocal on/off + language

    for p in INPUTS.values():
        p.touch(exist_ok=True)
    seen = {name: (p.stat().st_mtime_ns, p.read_text(encoding="utf-8", errors="ignore"))
            for name, p in INPUTS.items()}
    write_atomic(OUTPUT, "")  # start with a clear screen
    last_caption_time = 0.0
    last_sync_try = 0.0
    last_mic_check = 0.0
    started = time.time()
    obs_seen = False
    stop = False
    fragment, fragment_time = "", 0.0  # half sentence waiting for the rest

    def _stop(*_):
        nonlocal stop
        stop = True

    signal.signal(signal.SIGTERM, _stop)
    signal.signal(signal.SIGINT, _stop)
    stop_file = DATA / "stop_translator"
    stop_file.unlink(missing_ok=True)
    alive_file = DATA / "translator_alive"

    def show(raw: str, source: str, final: bool, display: bool = True) -> None:
        nonlocal last_caption_time
        t0 = time.time()
        caption = pipeline.process(raw, source)
        if not caption:
            return
        if display:
            write_atomic(OUTPUT, caption)
            last_caption_time = time.time()
        if final:
            log_transcript(source, raw, caption)
            info = pipeline.last
            pipeline.recorder.add(pipeline.mode, source, info.get("lang", ""),
                                  info.get("transcript", raw), caption)
            if app.gdocs is not None:
                app.gdocs.add(dt.datetime.now().strftime("%H:%M:%S"),
                              LANG_NAMES.get(info.get("lang", ""), info.get("lang", "")),
                              info.get("transcript", raw), caption)
            log.info("(%s %.2fs) %s  ->  %s", source, time.time() - t0, raw, caption)

    while not stop:
        now = time.time()
        if stop_file.exists() or app.quit:
            break
        # OBS closed (or crashed): its script stops updating obs_state.json.
        # (OBS also pauses the script for a few seconds while loading a Whisper
        # model, so allow a generous gap before deciding OBS is gone.)
        # Keep going while captions are being written into a Google Doc.
        alive = obs_bridge.state(max_age=75) is not None
        obs_seen = obs_seen or alive
        if with_obs:
            if (not alive and (obs_seen or now - started > 120) and not app.gdocs.active()
                    and now - app.last_client > 30):  # the app window may still be open
                log.info("OBS is no longer running: stopping")
                break
        elif (not alive and not app.gdocs.active() and now - app.last_client > STANDALONE_IDLE_QUIT
              and now - started > STANDALONE_IDLE_QUIT):
            # Window closed, no OBS, no Google Doc: don't keep listening to the mic unseen.
            log.info("Live Captions window closed: stopping")
            break
        mode = read_mode()  # also allows switching with --mode
        if mode != pipeline.mode:
            app.set_mode(mode)
        if last_caption_time and now - last_caption_time > CLEAR_AFTER_SECONDS:
            write_atomic(OUTPUT, "")  # don't leave the last words on screen
            last_caption_time = 0.0
        # Started before the OBS script was ready: sync again once it is.
        if (app.status.startswith(("no OBS", "OBS did not respond", "could not", "no Mic"))
                and now - last_sync_try > 3):
            last_sync_try = now
            if obs_bridge.state() is not None or app.status != "no OBS · " + app.speech.engine_status():
                app.sync()
        if now - last_mic_check > 2:  # Qwen3 follows the mic / mute chosen in OBS
            last_mic_check = now
            write_atomic(alive_file, str(int(now)))  # heartbeat for the OBS script's watchdog
            pipeline.vocab.reload_if_changed()  # edits to the vocabulary sheet apply live
            app.configure_speech()
            app.speech.follow_obs_microphone()
        on = captions_on()
        # Qwen3 captions
        pipeline.lenient = True
        if fragment and now - fragment_time > FRAGMENT_JOIN_SECONDS:
            try:  # nothing followed the half sentence: save it as it is
                show(fragment, "speech", True, display=False)
            except Exception:
                log.exception("Failed to process caption: %s", fragment)
            fragment = ""
        while True:
            try:
                text, final = app.speech.texts.get_nowait()
            except queue.Empty:
                break
            if not on:
                continue
            # A speaker's short pause cuts a sentence ("हर पल में" ... "निरंतर चल रहा है।"):
            # translate the half sentence together with what follows, so the
            # meaning isn't lost, and save only the joined line.
            if fragment and time.time() - fragment_time <= FRAGMENT_JOIN_SECONDS:
                text = fragment + " " + text
            if final:
                fragment = ""
                if is_fragment(text):
                    fragment, fragment_time = text, time.time()
            try:
                show(text, "speech", final and not fragment)
            except Exception:
                log.exception("Failed to process caption: %s", text)
        pipeline.lenient = False
        # Whisper (LocalVocal) captions, until Qwen3 is ready
        for name, path in INPUTS.items():
            try:
                mtime = path.stat().st_mtime_ns
            except FileNotFoundError:
                continue
            if mtime == seen[name][0]:
                continue
            raw = path.read_text(encoding="utf-8", errors="ignore").strip()
            changed = raw != seen[name][1]
            seen[name] = (mtime, raw)
            if not changed or not raw or not on or app.speech.engine == "qwen":
                continue
            try:
                show(raw, name, True)
            except Exception:
                log.exception("Failed to process caption: %s", raw)
        time.sleep(POLL_SECONDS)
    if app.speech.qwen is not None:
        app.speech.qwen.stop()
    write_atomic(OUTPUT, "")
    paths.PID_FILE.unlink(missing_ok=True)
    log.info("Translator stopped")


def self_test() -> None:
    pipeline = CaptionPipeline(Translator())
    samples = [
        ("Welcome to the satsang. Let us all do sewa with love.", "speech"),
        ("Dhan Nirankar Ji. Satguru Mata Sudiksha Ji Maharaj has blessed the sangat.", "speech"),
        ("This damn microphone is not working.", "speech"),
        ("आज हम सब मिलकर सत्संग में सेवा करेंगे।", "speech"),
        ("निरंकार हम सब पर कृपा करें। सिमरन करते रहें।", "speech"),
        ("सतगुरु माता सुदीक्षा जी महाराज का आशीर्वाद सबको मिले।", "speech"),
        ("गुरु जी ने ज्ञान दिया और माया से मुक्ति दिलाई।", "speech"),
        ("वह हरामी आदमी सत्संग में नहीं आया।", "speech"),
        ("तू ही निरंकार मैं तेरी शरण हाँ मैनूं बख्श लो", "speech"),
        ("Let us all say the Jaikara together.", "speech"),
        ("aaj hum sab milkar satsang mein sewa karenge", "speech"),
        ("Satguru Mata ji ka aashirwad sab ko mile", "speech"),
        ("kal ki meeting mein sab log time pe aana", "speech"),
        ("धन निरंकार जी।", "speech"),
        ("सत्कुरू माता जी का आशीरवाद सबको मिले।", "speech"),
        ("Açın. Açın", "speech"),
        ("Erne Dopenbahn", "speech"),
        ("Maraccio da ne?", "speech"),
        ("I'm not a good one", "speech"),
        ("きあえん キャージ", "speech"),
        ("-", "speech"),
    ]
    for text, source in samples:
        t0 = time.time()
        print(f"IN : {text}\nOUT: {pipeline.process(text, source)}   ({time.time() - t0:.2f}s)\n")


def self_test_modes() -> None:
    """Every sample in every caption mode."""
    pipeline = CaptionPipeline(Translator())
    samples = ["Welcome everyone to today's satsang. Let us do sewa together.",
               "Dhan Nirankar Ji. Satguru Mata Sudiksha Ji Maharaj has blessed the sangat.",
               "आज हम सब मिलकर सत्संग में सेवा करेंगे।",
               "सतगुरु माता सुदीक्षा जी महाराज का आशीर्वाद सबको मिले।",
               "aaj hum sab milkar satsang mein sewa karenge",
               "hamein hamesha pyaar se bolna chahiye"]
    for mode in MODES:
        pipeline.mode = mode
        print(f"===== {MODES[mode]['label']}")
        for s in samples:
            t0 = time.time()
            out = pipeline.process(s, "speech")
            print(f"  [{pipeline.last.get('lang', '-'):8}] {s}\n             -> {out}   ({time.time() - t0:.2f}s)")


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(prog="livecaptions", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--with-obs", action="store_true",
                    help="started by the OBS script: stop when OBS closes")
    ap.add_argument("--no-window", action="store_true", help="don't open the app window")
    ap.add_argument("--console", action="store_true", help="log to the terminal instead of translator.log")
    ap.add_argument("--test", action="store_true", help="translate sample sentences and exit")
    ap.add_argument("--test-modes", action="store_true", help="translate samples in every caption mode and exit")
    ap.add_argument("--download", action="store_true", help="download all models now and exit")
    ap.add_argument("--mode", choices=list(MODES),
                    help="switch caption direction (" + ", ".join(MODES) + ") of the running translator and exit")
    args = ap.parse_args(argv)
    DATA.mkdir(parents=True, exist_ok=True)
    handlers: list[logging.Handler] = []
    if not args.console and not (args.test or args.test_modes or args.download or args.mode):
        from logging.handlers import RotatingFileHandler
        handlers.append(RotatingFileHandler(paths.LOG_FILE, maxBytes=2_000_000, backupCount=2, encoding="utf-8"))
    else:
        handlers.append(logging.StreamHandler(sys.stdout))
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(name)s %(message)s", datefmt="%H:%M:%S",
                        handlers=handlers)
    if args.mode:
        write_atomic(MODE_FILE, args.mode)
        print(f"Caption mode set to: {MODES[args.mode]['label']}")
    elif args.download:
        s = models.Setup()
        s.start()
        while s.running:
            st = s.status()
            print(f"\r{st['percent']:3d}%  {st['message']:60.60}", end="", flush=True)
            time.sleep(1)
        print("\n" + (s.error or "All models are ready"))
    elif args.test:
        self_test()
    elif args.test_modes:
        self_test_modes()
    else:
        try:
            run(with_obs=args.with_obs, open_window=not args.with_obs and not args.no_window)
        except Exception:
            log.exception("Translator crashed")
            raise


if __name__ == "__main__":
    main()
